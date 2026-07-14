#!/usr/bin/env python3

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import (
    WD_ALIGN_PARAGRAPH,
    WD_BREAK,
    WD_TAB_ALIGNMENT,
    WD_TAB_LEADER,
)
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


REPORT = "实验二-TurtleBot4自主导航仿真-实验报告.docx"

FIGURE_CAPTIONS = [
    "TurtleBot4 仿真自主导航系统架构",
    "针对建图范围和路径长度反馈的改进前后对比",
    "初版短程 SLAM 证据基线",
    "达到 5.00 m / 58,506 cells 覆盖率门槛后保存的扩展地图",
    "AMCL、LaserScan 与长距离导航自动验收结果",
    "长距离目标发送前的 RViz：地图、机器人、LaserScan 与 AMCL 粒子",
    "长距离 A* 全局路径跨越多个 warehouse 通道",
    "长距离导航终点：RViz 显示 Feedback: reached",
]

TABLE_CAPTIONS = [
    "实验平台配置",
    "仿真代码包主要文件",
    "自定义 A* 与 Nav2 关键参数",
    "助教复验版本的自动验收门槛",
    "扩展 SLAM 覆盖率量化结果",
    "定位和激光稳定性结果",
    "长距离全局路径和 action 结果",
    "最终完整视频技术信息",
]


def add_page_break(paragraph):
    paragraph.add_run().add_break(WD_BREAK.PAGE)


def set_run_font(run, name, size=None, bold=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def insert_paragraph_after(paragraph, text="", style=None):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    from docx.text.paragraph import Paragraph

    inserted = Paragraph(new_p, paragraph._parent)
    if text:
        inserted.add_run(text)
    if style:
        inserted.style = style
    return inserted


def enable_field_updates(document):
    settings = document.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def style_cover(document):
    cover_values = {
        "组号：": "组号：________________",
        "课程名称：": "课程名称：机器人与智能系统综合实践",
        "姓名：": "姓名：________________",
        "院系：": "院系：________________",
        "专业：": "专业：________________",
        "学号：": "学号：________________",
        "指导老师：": "指导老师：________________",
        "实验平台：": "实验平台：ROS 2 Humble / TurtleBot4",
    }
    cover_title_map = {
        "本科实验报告": 26,
        "TurtleBot4 移动机器人自主导航与避障": 20,
        "Experiment 2: Autonomous Navigation of TurtleBot4": 14,
    }
    date_paragraph = None
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text in cover_title_map:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.space_after = Pt(12)
            for run in paragraph.runs:
                set_run_font(
                    run,
                    "黑体" if not text.startswith("Experiment") else "Times New Roman",
                    cover_title_map[text],
                    bold=True,
                )
        elif text.startswith(("组号：", "课程名称：", "姓", "院", "专", "学", "指导老师：", "实验平台：")):
            if text in cover_values:
                paragraph.text = cover_values[text]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.space_after = Pt(6)
            for run in paragraph.runs:
                set_run_font(run, "宋体", 12)
        elif text.startswith("2026 年 7 月"):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Cm(0)
            for run in paragraph.runs:
                set_run_font(run, "宋体", 12)
            date_paragraph = paragraph
            break
    if date_paragraph is not None:
        add_page_break(date_paragraph)
    section = document.sections[0]
    section.different_first_page_header_footer = True
    section.first_page_header.paragraphs[0].text = ""
    section.first_page_footer.paragraphs[0].text = ""


def prepare_bibliography(document):
    paragraphs = document.paragraphs
    start = next(
        index
        for index, paragraph in enumerate(paragraphs)
        if paragraph.text.strip().startswith("99 Clearpath Robotics")
    )
    paragraphs[start].insert_paragraph_before("参考文献", "Heading 1")
    for number, paragraph in enumerate(paragraphs[start:start + 7], start=1):
        if number == 1:
            for run in paragraph.runs:
                if run.text.strip() == "99":
                    run.text = ""
        for run in paragraph.runs:
            if run.text:
                run.text = f"[{number}] {run.text}"
                break
        paragraph.paragraph_format.left_indent = Cm(0.74)
        paragraph.paragraph_format.first_line_indent = Cm(-0.74)


def number_headings(document):
    chapter = 0
    subsection = 0
    current = ""
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if paragraph.style.name == "Heading 1" and text not in {"摘要", "目录"}:
            subsection = 0
            if text == "仿真复现命令":
                current = "A"
                prefix = "附录 A"
            elif text == "报告生成方法":
                current = "B"
                prefix = "附录 B"
            else:
                chapter += 1
                current = str(chapter)
                prefix = current
            paragraph.text = f"{prefix}  {text}"
        elif paragraph.style.name == "Heading 2" and current:
            subsection += 1
            paragraph.text = f"{current}.{subsection}  {text}"


def insert_contents_page(document):
    keywords = next(
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text.strip().startswith("关键词：")
    )
    add_page_break(keywords)
    title = insert_paragraph_after(keywords, "目录", "Heading 1")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.first_line_indent = Cm(0)
    entries = [
        (1, "1  实验目的与要求", 7),
        (2, "1.1  实验目的", 7),
        (2, "1.2  仿真实验要求", 7),
        (2, "1.3  实验范围说明", 7),
        (1, "2  实验平台与 TurtleBot4 工作原理", 8),
        (2, "2.1  软硬件与仿真平台", 8),
        (2, "2.2  TurtleBot4 主要硬件构成", 8),
        (2, "2.3  远程主机连接原理", 9),
        (1, "3  自主导航系统设计", 9),
        (2, "3.1  总体架构", 9),
        (2, "3.2  SLAM Toolbox 增量建图", 10),
        (2, "3.3  AMCL 自定位", 10),
        (2, "3.4  代价感知 A* 全局路径规划", 10),
        (2, "3.5  Nav2 闭环执行", 11),
        (1, "4  软件实现", 11),
        (2, "4.1  代码包结构", 11),
        (2, "4.2  仿真启动模式", 12),
        (2, "4.3  关键导航参数", 13),
        (2, "4.4  定位与长距离验收门槛", 13),
        (1, "5  实验步骤", 13),
        (2, "5.1  编译与静态验证", 13),
        (2, "5.2  扩展 SLAM 建图", 14),
        (2, "5.3  稳定 AMCL 初始化与规划器检查", 14),
        (2, "5.4  发送长距离导航目标", 15),
        (1, "6  实验结果与分析", 16),
        (2, "6.1  改进前问题量化与验收策略", 16),
        (2, "6.2  扩展 SLAM 覆盖率结果", 16),
        (2, "6.3  AMCL、TF 与 LaserScan 稳定性", 17),
        (2, "6.4  长距离 A* 路径与闭环执行", 19),
        (2, "6.5  完整无加速视频验收", 20),
        (2, "6.6  仿真时钟与低负载配置分析", 21),
        (1, "7  问题诊断与改进", 21),
        (2, "7.1  启动参数作用域", 21),
        (2, "7.2  RViz 视角确定性", 22),
        (2, "7.3  慢速主机的 lifecycle 启动顺序", 22),
        (2, "7.4  建图范围与长路径的自动门槛", 22),
        (2, "7.5  雷达云漂移的量化判定", 22),
        (2, "7.6  完整视频证据保真", 22),
        (2, "7.7  后续优化方向", 23),
        (1, "8  结论", 23),
        (1, "9  参考文献", 24),
        (1, "附录 A  仿真复现命令", 25),
        (1, "附录 B  报告生成方法", 26),
    ]
    previous = title
    for level, name, page in entries:
        entry = insert_paragraph_after(previous, f"{name}\t{page}")
        entry.paragraph_format.left_indent = Cm(0.74 if level == 2 else 0)
        entry.paragraph_format.first_line_indent = Cm(0)
        entry.paragraph_format.space_before = Pt(0)
        entry.paragraph_format.space_after = Pt(0)
        entry.paragraph_format.line_spacing = 1.0
        entry.paragraph_format.tab_stops.add_tab_stop(
            Cm(15.0), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS
        )
        for run in entry.runs:
            set_run_font(run, "宋体", 9.5, bold=level == 1)
        previous = entry
    add_page_break(previous)


def number_captions(document):
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text in FIGURE_CAPTIONS:
            number = FIGURE_CAPTIONS.index(text) + 1
            paragraph.text = f"图 {number}  {text}"
            paragraph.style = "Caption"
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Cm(0)
        elif text in TABLE_CAPTIONS:
            number = TABLE_CAPTIONS.index(text) + 1
            paragraph.text = f"表 {number}  {text}"
            paragraph.style = "Caption"
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.keep_with_next = True
            if number in {1, 4, 5, 6}:
                paragraph.paragraph_format.page_break_before = True


def style_body(document):
    for paragraph in document.paragraphs:
        if paragraph.style.name == "Source Code":
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.space_before = Pt(4)
            paragraph.paragraph_format.space_after = Pt(4)
            for run in paragraph.runs:
                set_run_font(run, "Consolas", 9)
        elif paragraph.style.name.startswith("Heading"):
            paragraph.paragraph_format.first_line_indent = Cm(0)
        elif paragraph.style.name == "Caption":
            paragraph.paragraph_format.first_line_indent = Cm(0)
        elif paragraph.text.strip() and paragraph.text.strip() not in {
            "本科实验报告",
            "TurtleBot4 移动机器人自主导航与避障",
            "Experiment 2: Autonomous Navigation of TurtleBot4",
            "目录",
        }:
            if paragraph.alignment != WD_ALIGN_PARAGRAPH.CENTER:
                paragraph.paragraph_format.line_spacing = 1.5

    for shape in document.inline_shapes:
        parent = shape._inline.getparent().getparent()
        from docx.text.paragraph import Paragraph

        Paragraph(parent, document._body).alignment = WD_ALIGN_PARAGRAPH.CENTER


def style_tables(document):
    for table in document.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = "Table Grid"
        for row_index, row in enumerate(table.rows):
            tr_properties = row._tr.get_or_add_trPr()
            if tr_properties.find(qn("w:cantSplit")) is None:
                tr_properties.append(OxmlElement("w:cantSplit"))
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.first_line_indent = Cm(0)
                    paragraph.paragraph_format.space_after = Pt(0)
                    paragraph.paragraph_format.line_spacing = 1.15
                    paragraph.paragraph_format.keep_with_next = (
                        row_index < len(table.rows) - 1
                    )
                    if row_index == 0:
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in paragraph.runs:
                        set_run_font(run, "宋体", 8.5, bold=row_index == 0)


def main():
    document = Document(REPORT)
    prepare_bibliography(document)
    number_headings(document)
    style_cover(document)
    insert_contents_page(document)
    number_captions(document)
    style_body(document)
    style_tables(document)
    enable_field_updates(document)
    document.core_properties.title = "TurtleBot4 移动机器人自主导航与避障"
    document.core_properties.subject = "机器人与智能系统综合实践·实验二"
    document.core_properties.keywords = "TurtleBot4, ROS 2, Gazebo, SLAM, AMCL, Nav2, A*"
    document.save(REPORT)


if __name__ == "__main__":
    main()
