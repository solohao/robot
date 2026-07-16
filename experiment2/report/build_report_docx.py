#!/usr/bin/env python3

import argparse
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from lxml import etree


ROOT = Path(__file__).resolve().parent
TEMPLATE = ROOT / "template" / "实验报告模板.docx"
BODY_HEADING = "一、实验目的和要求（必填）"

NAMESPACES = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
}


def set_east_asia_font(run_or_style, font_name):
    run_or_style.font.name = "Times New Roman"
    run_or_style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font_name)


def set_table_borders(table, color="FFFFFF", size="0"):
    table_properties = table._tbl.tblPr
    borders = table_properties.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        table_properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "nil" if size == "0" else "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_bottom_border(cell, color="666666", size="6"):
    cell_properties = cell._tc.get_or_add_tcPr()
    borders = cell_properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        cell_properties.append(borders)
    bottom = borders.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        borders.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:color"), color)


def set_repeat_table_header(row):
    row_properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    row_properties.append(repeat)


def set_update_fields(document):
    settings = document.settings._element
    existing = settings.find(qn("w:updateFields"))
    if existing is None:
        existing = OxmlElement("w:updateFields")
        settings.append(existing)
    existing.set(qn("w:val"), "true")


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run("第 ")
    set_east_asia_font(run, "宋体")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, end])
    suffix = paragraph.add_run(" 页")
    set_east_asia_font(suffix, "宋体")


def read_logo():
    with ZipFile(TEMPLATE) as package:
        document = etree.fromstring(package.read("word/document.xml"))
        relationships = etree.fromstring(
            package.read("word/_rels/document.xml.rels")
        )
        relationship_by_id = {
            relationship.get("Id"): relationship.get("Target")
            for relationship in relationships
        }
        first_image = document.find(".//a:blip", namespaces=NAMESPACES)
        if first_image is None:
            raise ValueError("template logo not found")
        relationship_id = first_image.get(f"{{{NAMESPACES['r']}}}embed")
        target = relationship_by_id[relationship_id]
        return package.read(f"word/{target}")


def element_text(element):
    text_tags = {qn("w:t"), qn("m:t")}
    return "".join(
        text.text or "" for text in element.iter() if text.tag in text_tags
    ).strip()


def retain_report_body(document):
    body = document._element.body
    elements = list(body)
    start = None
    for index, element in enumerate(elements):
        if element.tag == qn("w:p") and element_text(element).startswith(BODY_HEADING):
            start = index
            break
    if start is None:
        raise ValueError(f"could not find body heading {BODY_HEADING!r}")

    retained = [
        deepcopy(element)
        for element in elements[start:]
        if element.tag != qn("w:sectPr")
    ]
    for element in list(body):
        if element.tag != qn("w:sectPr"):
            body.remove(element)
    return retained


def configure_document(document):
    template = Document(TEMPLATE)
    source_section = template.sections[0]
    section = document.sections[0]
    section.page_width = source_section.page_width
    section.page_height = source_section.page_height
    section.top_margin = source_section.top_margin
    section.right_margin = source_section.right_margin
    section.bottom_margin = source_section.bottom_margin
    section.left_margin = source_section.left_margin
    section.header_distance = source_section.header_distance
    section.footer_distance = source_section.footer_distance

    normal = document.styles["Normal"]
    set_east_asia_font(normal, "宋体")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.first_line_indent = Cm(0.74)
    normal.paragraph_format.space_after = Pt(0)

    for name, size, bold in (
        ("Title", 18, True),
        ("Heading 1", 14, True),
        ("Heading 2", 12, True),
        ("Heading 3", 10.5, True),
    ):
        if name not in document.styles:
            continue
        style = document.styles[name]
        set_east_asia_font(style, "黑体")
        style.font.size = Pt(size)
        style.font.bold = bold
        style.paragraph_format.first_line_indent = Pt(0)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(4)

    for style_name in ("Caption", "Image Caption", "Table Caption"):
        if style_name in document.styles:
            style = document.styles[style_name]
            set_east_asia_font(style, "宋体")
            style.font.size = Pt(9)
            style.font.bold = False

    footer = section.footer
    footer.is_linked_to_previous = False
    footer_paragraph = footer.paragraphs[0]
    footer_paragraph.clear()
    add_page_number(footer_paragraph)

    set_update_fields(document)
    properties = document.core_properties
    properties.title = "实验二：TurtleBot4 移动机器人自主导航与避障"
    properties.subject = "计算机仿真报告与实地实验填写模板"
    properties.author = ""
    properties.keywords = "ROS 2, TurtleBot4, Nav2, A*, AMCL, SLAM"


def add_cover(document, logo):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.space_after = Pt(46)
    run = paragraph.add_run("组号：________________")
    set_east_asia_font(run, "宋体")
    run.font.size = Pt(10.5)

    logo_paragraph = document.add_paragraph()
    logo_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    logo_paragraph.paragraph_format.first_line_indent = Pt(0)
    logo_paragraph.paragraph_format.space_after = Pt(24)
    logo_paragraph.add_run().add_picture(BytesIO(logo), width=Cm(6.2))

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.first_line_indent = Pt(0)
    title.paragraph_format.space_after = Pt(58)
    title_run = title.add_run("本科实验报告")
    set_east_asia_font(title_run, "黑体")
    title_run.font.size = Pt(18)
    title_run.font.bold = True

    fields = (
        ("课程名称：", "机器人与智能系统综合实践"),
        ("姓　　名：", ""),
        ("院　　系：", ""),
        ("专　　业：", ""),
        ("学　　号：", ""),
        ("指导老师：", ""),
        ("选课时间：", ""),
    )
    table = document.add_table(rows=len(fields), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    for row, (label, value) in zip(table.rows, fields):
        row.height = Cm(0.82)
        row.cells[0].width = Cm(3.1)
        row.cells[1].width = Cm(7.2)
        row.cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        row.cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        label_paragraph = row.cells[0].paragraphs[0]
        label_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        label_paragraph.paragraph_format.first_line_indent = Pt(0)
        label_run = label_paragraph.add_run(label)
        set_east_asia_font(label_run, "宋体")
        label_run.font.size = Pt(10.5)
        value_paragraph = row.cells[1].paragraphs[0]
        value_paragraph.paragraph_format.first_line_indent = Pt(0)
        value_run = value_paragraph.add_run(value)
        set_east_asia_font(value_run, "宋体")
        value_run.font.size = Pt(10.5)
        set_cell_bottom_border(row.cells[1])

    date = document.add_paragraph()
    date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date.paragraph_format.first_line_indent = Pt(0)
    date.paragraph_format.space_before = Pt(34)
    date_run = date.add_run("________ 年 ____ 月 ____ 日")
    set_east_asia_font(date_run, "宋体")
    date_run.font.size = Pt(10.5)
    document.add_page_break()


def add_second_page_header(document, logo):
    header = document.add_table(rows=1, cols=2)
    header.alignment = WD_TABLE_ALIGNMENT.CENTER
    header.autofit = False
    set_table_borders(header)
    left, right = header.rows[0].cells
    left.width = Cm(9.4)
    right.width = Cm(5.2)
    left.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP

    left_paragraph = left.paragraphs[0]
    left_paragraph.paragraph_format.first_line_indent = Pt(0)
    left_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    left_paragraph.add_run().add_picture(BytesIO(logo), width=Cm(3.7))
    title = left_paragraph.add_run(" 实验报告")
    set_east_asia_font(title, "黑体")
    title.font.size = Pt(21)
    title.font.bold = True

    for index, label in enumerate(("专业", "姓名", "学号", "日期", "地点")):
        paragraph = right.paragraphs[0] if index == 0 else right.add_paragraph()
        paragraph.paragraph_format.first_line_indent = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = paragraph.add_run(f"{label}：________________")
        set_east_asia_font(run, "宋体")
        run.font.size = Pt(9)

    information = document.add_table(rows=3, cols=1)
    information.alignment = WD_TABLE_ALIGNMENT.CENTER
    information.autofit = False
    set_table_borders(information)
    lines = (
        "课程名称：机器人与智能系统综合实践　　指导老师：____________　　成绩：________",
        "实验名称：TurtleBot4 移动机器人自主导航与避障",
        "实验类型：综合性实验（计算机仿真 + 实地实验）　　同组学生姓名：________________",
    )
    for cell, line in zip((row.cells[0] for row in information.rows), lines):
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.first_line_indent = Pt(0)
        paragraph.paragraph_format.space_after = Pt(1)
        run = paragraph.add_run(line)
        set_east_asia_font(run, "宋体")
        run.font.size = Pt(9)

    rule = document.add_paragraph()
    rule.paragraph_format.first_line_indent = Pt(0)
    rule.paragraph_format.space_before = Pt(0)
    rule.paragraph_format.space_after = Pt(4)
    rule_run = rule.add_run("—" * 49)
    set_east_asia_font(rule_run, "宋体")
    rule_run.font.size = Pt(8)

    checklist = document.add_table(rows=4, cols=2)
    checklist.alignment = WD_TABLE_ALIGNMENT.CENTER
    checklist.autofit = False
    set_table_borders(checklist)
    entries = (
        ("一、实验目的和要求（必填）", "二、实验内容和原理（必填）"),
        ("三、主要仪器设备（必填）", "四、操作方法和实验步骤"),
        ("五、实验数据记录和处理", "六、实验结果与分析（必填）"),
        ("七、讨论、心得", ""),
    )
    for row, values in zip(checklist.rows, entries):
        for cell, text in zip(row.cells, values):
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.first_line_indent = Pt(0)
            paragraph.paragraph_format.space_after = Pt(1)
            run = paragraph.add_run(text)
            set_east_asia_font(run, "宋体")
            run.font.size = Pt(9)

    separator = document.add_paragraph()
    separator.paragraph_format.first_line_indent = Pt(0)
    separator.paragraph_format.space_before = Pt(2)
    separator.paragraph_format.space_after = Pt(2)
    separator_run = separator.add_run("—" * 49)
    set_east_asia_font(separator_run, "宋体")
    separator_run.font.size = Pt(8)


def repair_gridless_tables(document):
    section = document.sections[0]
    available_width = int(
        section.page_width.twips
        - section.left_margin.twips
        - section.right_margin.twips
    )
    fractions = {
        2: (0.34, 0.66),
        3: (0.36, 0.32, 0.32),
        4: (0.38, 0.21, 0.22, 0.19),
    }

    gridless_tables = [
        table
        for table in document._element.body.iter(qn("w:tbl"))
        if (
            table.find(qn("w:tblGrid")) is None
            or not len(table.find(qn("w:tblGrid")))
        )
    ]
    for table in gridless_tables:
        source_rows = table.findall(qn("w:tr"))
        source_cells = [
            [element_text(cell) for cell in row.findall(qn("w:tc"))]
            for row in source_rows
        ]
        column_count = max((len(row) for row in source_cells), default=0)
        if not source_cells or column_count == 0:
            table.getparent().remove(table)
            continue
        column_fractions = fractions.get(
            column_count,
            tuple(1.0 / column_count for _ in range(column_count)),
        )
        widths = [int(available_width * fraction) for fraction in column_fractions]
        widths[-1] += available_width - sum(widths)

        rebuilt = document.add_table(
            rows=len(source_cells),
            cols=column_count,
        )
        rebuilt.style = "Table Grid"
        rebuilt.alignment = WD_TABLE_ALIGNMENT.CENTER
        rebuilt.autofit = False
        for row_index, values in enumerate(source_cells):
            for column_index in range(column_count):
                cell = rebuilt.rows[row_index].cells[column_index]
                cell.width = Pt(widths[column_index] / 20.0)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
                paragraph = cell.paragraphs[0]
                paragraph.paragraph_format.first_line_indent = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                run = paragraph.add_run(
                    values[column_index] if column_index < len(values) else ""
                )
                set_east_asia_font(run, "宋体")
                run.font.size = Pt(9)
                run.font.bold = row_index == 0

        table.addprevious(rebuilt._tbl)
        table.getparent().remove(table)


def normalize_body(document):
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text.startswith(("三、", "五、", "六、")):
            paragraph.paragraph_format.page_break_before = True
        if text.startswith("图 ") or text.startswith("表 "):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Pt(0)
        if text.startswith("附录 ") or text == "参考文献":
            paragraph.paragraph_format.keep_with_next = True

    for table in document.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        if table.rows:
            set_repeat_table_header(table.rows[0])
        for row_index, row in enumerate(table.rows):
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.first_line_indent = Pt(0)
                    paragraph.paragraph_format.space_after = Pt(0)
                    for run in paragraph.runs:
                        set_east_asia_font(run, "宋体")
                        run.font.size = Pt(9)
                        if row_index == 0:
                            run.font.bold = True


def build(input_path, output_path):
    document = Document(input_path)
    retained = retain_report_body(document)
    configure_document(document)
    logo = read_logo()
    add_cover(document, logo)
    add_second_page_header(document, logo)

    body = document._element.body
    section_properties = body.sectPr
    for element in retained:
        body.insert(body.index(section_properties), element)

    repair_gridless_tables(document)
    normalize_body(document)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    arguments = parser.parse_args()
    build(arguments.input, arguments.output)


if __name__ == "__main__":
    main()
