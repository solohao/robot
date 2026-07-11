#!/usr/bin/env python3

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT = "reference.docx"


def set_east_asia_font(style, font_name):
    style.font.name = font_name
    style._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, end])


def main():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.6)
    section.header_distance = Cm(1.2)
    section.footer_distance = Cm(1.2)

    normal = doc.styles["Normal"]
    set_east_asia_font(normal, "宋体")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.first_line_indent = Cm(0.74)

    title = doc.styles["Title"]
    set_east_asia_font(title, "黑体")
    title.font.size = Pt(22)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0x11, 0x18, 0x27)
    title.paragraph_format.space_after = Pt(18)

    subtitle = doc.styles["Subtitle"]
    set_east_asia_font(subtitle, "黑体")
    subtitle.font.size = Pt(14)
    subtitle.font.color.rgb = RGBColor(0x37, 0x41, 0x51)

    for name, size, before, after in (
        ("Heading 1", 16, 18, 12),
        ("Heading 2", 14, 14, 8),
        ("Heading 3", 12, 10, 6),
    ):
        style = doc.styles[name]
        set_east_asia_font(style, "黑体")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0x11, 0x18, 0x27)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    caption = doc.styles["Caption"]
    set_east_asia_font(caption, "宋体")
    caption.font.size = Pt(9)
    caption.font.color.rgb = RGBColor(0x37, 0x41, 0x51)
    caption.paragraph_format.first_line_indent = Cm(0)

    if "Abstract" not in doc.styles:
        abstract = doc.styles.add_style("Abstract", WD_STYLE_TYPE.PARAGRAPH)
    else:
        abstract = doc.styles["Abstract"]
    set_east_asia_font(abstract, "宋体")
    abstract.font.size = Pt(10.5)
    abstract.paragraph_format.line_spacing = 1.5

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header_run = header.add_run("机器人与智能系统综合实践 · 实验二")
    header_run.font.size = Pt(9)
    header_run.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
    header_run.font.name = "宋体"
    header_run._element.get_or_add_rPr().get_or_add_rFonts().set(
        qn("w:eastAsia"), "宋体"
    )
    add_page_field(section.footer.paragraphs[0])

    doc.add_paragraph("样式参考文档")
    doc.save(OUTPUT)


if __name__ == "__main__":
    main()
